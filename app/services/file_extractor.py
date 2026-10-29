from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".csv", ".txt"}

_SKIP_DIR_NAMES = {"__MACOSX"}


@dataclass(frozen=True)
class ExtractedFile:
    path: Path
    display_name: str
    fmt: str
    text: str
    status: str
    error: str | None = None

    @property
    def char_count(self) -> int:
        return len(self.text)


def discover_files(folder: Path) -> list[Path]:
    """Return every regular file under *folder* (recursive, deterministic order).

    Hidden/editor-temp artifacts and OS metadata folders are ignored. Unsupported
    formats are intentionally returned so the pipeline can report them explicitly.
    """
    if not folder.exists():
        raise FileNotFoundError(f"Input folder does not exist: {folder}")
    if not folder.is_dir():
        raise NotADirectoryError(f"Input path is not a folder: {folder}")

    files: list[Path] = []
    for path in sorted(folder.rglob("*")):
        if not path.is_file():
            continue
        if any(part in _SKIP_DIR_NAMES for part in path.parts):
            continue
        name = path.name
        if name.startswith(".") or name.startswith("~$"):
            continue
        files.append(path)
    return files


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip("\n")


def _extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:  # pragma: no cover - depends on file
            raise ValueError("encrypted PDF cannot be read") from exc
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = [p.text for p in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            parts.append(", ".join(cells))

    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                if paragraph.text.strip():
                    parts.append(paragraph.text)

    return "\n".join(part for part in parts if part is not None)


def _extract_csv(path: Path) -> str:
    raw = _decode(path.read_bytes())
    # Validate/normalize through the csv module so quoted commas survive; fall
    # back to raw text if the file is malformed.
    try:
        rows = list(csv.reader(io.StringIO(raw)))
    except csv.Error:
        return raw
    return "\n".join(",".join(row) for row in rows)


def _extract_txt(path: Path) -> str:
    return _decode(path.read_bytes())


_EXTRACTORS = {
    ".pdf": ("pdf", _extract_pdf),
    ".docx": ("docx", _extract_docx),
    ".csv": ("csv", _extract_csv),
    ".txt": ("txt", _extract_txt),
}


def extract_file(path: Path, root: Path) -> ExtractedFile:
    """Extract text from one file, reporting an explicit status on every path."""
    try:
        display_name = path.relative_to(root).as_posix()
    except ValueError:
        display_name = path.name

    suffix = path.suffix.lower()
    entry = _EXTRACTORS.get(suffix)
    if entry is None:
        return ExtractedFile(
            path=path,
            display_name=display_name,
            fmt=suffix.lstrip(".") or "unknown",
            text="",
            status="skipped",
            error=f"unsupported format '{suffix or 'no extension'}'",
        )

    fmt, extractor = entry
    try:
        text = normalize_text(extractor(path))
    except Exception as exc:
        return ExtractedFile(
            path=path,
            display_name=display_name,
            fmt=fmt,
            text="",
            status="failed",
            error=f"{type(exc).__name__}: {exc}"[:300],
        )

    if not text.strip():
        return ExtractedFile(
            path=path,
            display_name=display_name,
            fmt=fmt,
            text="",
            status="empty",
            error="no extractable text",
        )

    return ExtractedFile(
        path=path,
        display_name=display_name,
        fmt=fmt,
        text=text,
        status="processed",
    )


def extract_folder(folder: Path) -> list[ExtractedFile]:
    return [extract_file(path, folder) for path in discover_files(folder)]
