from __future__ import annotations

from pathlib import Path

from docx import Document

from app.services.file_extractor import discover_files, extract_file


def test_discovers_supported_and_unsupported_files(tmp_path: Path) -> None:
    (tmp_path / "a.txt").write_text("hello", encoding="utf-8")
    (tmp_path / "b.csv").write_text("name,value\nA,1\n", encoding="utf-8")
    (tmp_path / "c.bin").write_bytes(b"\x00\x01\x02")
    (tmp_path / "empty.txt").write_text("", encoding="utf-8")
    (tmp_path / ".hidden.txt").write_text("hidden", encoding="utf-8")

    names = [p.name for p in discover_files(tmp_path)]
    assert names == ["a.txt", "b.csv", "c.bin", "empty.txt"]

    assert extract_file(tmp_path / "a.txt", tmp_path).status == "processed"
    assert extract_file(tmp_path / "c.bin", tmp_path).status == "skipped"
    assert extract_file(tmp_path / "empty.txt", tmp_path).status == "empty"


def test_nested_files_and_display_names(tmp_path: Path) -> None:
    nested = tmp_path / "sub"
    nested.mkdir()
    (nested / "nested.txt").write_text("nested text", encoding="utf-8")

    extracted = extract_file(nested / "nested.txt", tmp_path)
    assert extracted.status == "processed"
    assert extracted.display_name == "sub/nested.txt"


def test_csv_preserves_quoted_fields(tmp_path: Path) -> None:
    path = tmp_path / "quoted.csv"
    path.write_text('name,note\n"Smith, John",ok\n', encoding="utf-8")
    text = extract_file(path, tmp_path).text
    assert "Smith, John" in text
    assert "ok" in text


def test_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "statement.docx"
    document = Document()
    document.add_paragraph("Statement for Test User")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Date"
    table.cell(0, 1).text = "Amount"
    document.save(str(path))

    extracted = extract_file(path, tmp_path)
    assert extracted.status == "processed"
    assert "Statement for Test User" in extracted.text
    assert "Date" in extracted.text and "Amount" in extracted.text
